# -*- coding: utf-8 -*-
"""slides.md から編集可能な .pptx を生成する。

    python scripts/build_pptx.py <slides.md> [<slides.md> ...]

Marpのレンダリングと同じ 1280x720px キャンバス(= 13.33x7.5in, 1px = 9525EMU)で
レイアウトし、Sports Force デザインシステムの配色を再現する。
テキストはすべてテキストボックス、図解は事前に diagram_export.py で変換したPNGを埋め込む。
講師ノート(<!-- ノート: ... -->)はスピーカーノートに入る。
"""
import re
import sys
from math import ceil
from pathlib import Path

from PIL import Image
from pygments import lex
from pygments.lexers import get_lexer_by_name
from pygments.lexers.special import TextLexer
from pygments.token import Comment, Keyword, Name, Number, String
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

PX = 9525  # EMU per px (96dpi)
SLIDE_W, SLIDE_H = 1280, 720
PAD_X = 76  # section padding
CONTENT_W = SLIDE_W - PAD_X * 2

# Sports Force カラー
INK = RGBColor(0x0E, 0x0E, 0x10)
TEXT = RGBColor(0x33, 0x33, 0x3A)
MUTED = RGBColor(0x8A, 0x8A, 0x93)
FAINT = RGBColor(0xB7, 0xB7, 0xBE)
PINK = RGBColor(0xE6, 0x2F, 0x9A)
CODE_PINK = RGBColor(0xB8, 0x1E, 0x72)
BG_SOFT = RGBColor(0xF7, 0xF7, 0xF9)
BORDER = RGBColor(0xDD, 0xDD, 0xE2)
LEAD_BG = RGBColor(0x14, 0x14, 0x18)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LEAD_SUB = RGBColor(0x9B, 0x9B, 0xA4)  # lead上の半透明白の近似

JP_FONT = "Noto Sans JP"
MONO_FONT = "Consolas"

GRAD_XML = (
    '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    "<a:gsLst>"
    '<a:gs pos="0"><a:srgbClr val="0A33FF"/></a:gs>'
    '<a:gs pos="34000"><a:srgbClr val="8330C7"/></a:gs>'
    '<a:gs pos="58000"><a:srgbClr val="E62F9A"/></a:gs>'
    '<a:gs pos="78000"><a:srgbClr val="FF2C61"/></a:gs>'
    '<a:gs pos="100000"><a:srgbClr val="FF2E0D"/></a:gs>'
    "</a:gsLst>"
    '<a:lin ang="0" scaled="1"/>'
    "</a:gradFill>"
)


def px(v):
    return Emu(round(v * PX))


def pt(size_px):
    return Pt(round(size_px * 0.75 * 2) / 2)


# ---------------------------------------------------------------- parsing


def parse_front_matter(text):
    """front-matterを取り除き、(headerの値, 本文) を返す。"""
    header = ""
    if text.startswith("---"):
        end = text.index("\n---", 3)
        fm = text[:end]
        m = re.search(r'^header:\s*"?([^"\n]+)"?\s*$', fm, re.M)
        if m:
            header = m.group(1)
        text = text[end + 4:]
    return header, text


def parse_slides(body):
    """本文をスライドごとの dict {cls, note, blocks} のリストに分解する。"""
    slides = []
    for chunk in re.split(r"\n---\n", body):
        cls = None
        notes = []
        for c in re.findall(r"<!--(.*?)-->", chunk, re.S):
            c = c.strip()
            m = re.match(r"_class:\s*(\w+)", c)
            if m:
                cls = m.group(1)
            else:
                notes.append(re.sub(r"^ノート:\s*", "", c))
        chunk = re.sub(r"<!--.*?-->", "", chunk, flags=re.S)
        slides.append({"cls": cls, "note": "\n".join(notes).strip(),
                       "blocks": parse_blocks(chunk)})
    return slides


def parse_blocks(text):
    blocks = []
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", lang, code))
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if m:
            wm = re.search(r"w:(\d+)", m.group(1))
            blocks.append(("img", m.group(2), int(wm.group(1)) if wm else 950))
            i += 1
            continue
        if stripped.startswith("# ") and not stripped.startswith("## "):
            # 連続するh1行は1ブロックにまとめる(leadのタイトル2行)
            h1 = [stripped[2:]]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith("# ") \
                    and not lines[i + 1].strip().startswith("## "):
                i += 1
                h1.append(lines[i].strip()[2:])
            blocks.append(("h1", h1))
            i += 1
            continue
        if stripped.startswith("## "):
            blocks.append(("h2", stripped[3:]))
            i += 1
            continue
        lm = re.match(r"^(\s*)(-|\*|\d+\.)\s+(.*)$", line)
        if lm:
            items = []
            while i < len(lines):
                lm = re.match(r"^(\s*)(-|\*|\d+\.)\s+(.*)$", lines[i])
                if not lm:
                    break
                level = min(len(lm.group(1)) // 2, 1)
                ordered = lm.group(2)[0].isdigit()
                items.append((level, ordered, lm.group(3)))
                i += 1
            blocks.append(("list", items))
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        # 段落: 連続する通常行をまとめる(Marpは改行を<br>にするので行を保持)
        para = [stripped]
        while i + 1 < len(lines):
            nxt = lines[i + 1].strip()
            if not nxt or nxt.startswith(("#", "-", "*", "|", "`", "!")) \
                    or re.match(r"^\d+\.\s", nxt):
                break
            i += 1
            para.append(nxt)
        blocks.append(("p", para))
        i += 1
    return blocks


def parse_inline(text, bold=False):
    """**bold** と `code` を run のリスト (text, bold, code) に分解する。"""
    runs = []
    for token in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            runs.extend(parse_inline(token[2:-2], bold=True))
        elif token.startswith("`") and token.endswith("`"):
            runs.append((token[1:-1], bold, True))
        else:
            runs.append((token, bold, False))
    return runs


def plain_text(text):
    return re.sub(r"\*\*(.+?)\*\*|`([^`]+)`", lambda m: m.group(1) or m.group(2), text)


# ------------------------------------------------------- syntax highlight

# GitHub Light系の配色(明るいグレー背景 #F7F7F9 上で読める色)
CODE_COLORS = [
    (Comment, MUTED),
    (String, RGBColor(0x0A, 0x30, 0x69)),
    (Number, RGBColor(0x05, 0x50, 0xAE)),
    (Keyword, RGBColor(0xCF, 0x22, 0x2E)),
    (Name.Function, RGBColor(0x82, 0x50, 0xDF)),
    (Name.Class, RGBColor(0x95, 0x38, 0x00)),
    (Name.Builtin, RGBColor(0x82, 0x50, 0xDF)),
    (Name.Decorator, RGBColor(0x82, 0x50, 0xDF)),
]


def code_color(ttype):
    for base, color in CODE_COLORS:
        if ttype in base:
            return color
    return TEXT


def highlight_lines(code_lines, lang):
    """コード行のリストを [(text, color), ...] の行リストに変換する。"""
    try:
        lexer = get_lexer_by_name(lang or "text")
    except Exception:
        lexer = TextLexer()
    lines = [[]]
    for ttype, value in lex("\n".join(code_lines), lexer):
        for k, part in enumerate(value.split("\n")):
            if k > 0:
                lines.append([])
            if part:
                lines[-1].append((part, code_color(ttype)))
    while len(lines) > len(code_lines) and not lines[-1]:
        lines.pop()
    return lines


# ------------------------------------------------------------ text metrics


def text_units(s):
    """全角=1.0 / 半角=0.52 でem単位の幅を見積もる。"""
    return sum(1.0 if ord(ch) >= 0x1100 else 0.52 for ch in s)


def est_lines(text, fs, avail):
    return max(1, ceil(text_units(text) * fs / avail))


def mono_units(s):
    """等幅フォント(Consolas)用の幅見積もり。半角は全角より広めの0.62em。"""
    return sum(1.0 if ord(ch) >= 0x1100 else 0.62 for ch in s)


# ------------------------------------------------------------ pptx helpers


def set_run(run, size_px, color=TEXT, bold=False, mono=False, spacing_pt=None):
    f = run.font
    f.size = pt(size_px)
    f.color.rgb = color
    # 常に太字で出力する: Google Slides等でNoto Sans JPが細字に代替されるため。
    # 強調(**)はboldではなくアクセント色(PINK)で区別される
    f.bold = True
    f.name = MONO_FONT if mono else JP_FONT
    rPr = run._r.get_or_add_rPr()
    ea = parse_xml(
        f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{JP_FONT}"/>'
    )
    rPr.append(ea)
    if spacing_pt:
        rPr.set("spc", str(int(spacing_pt * 100)))


def fill_runs(para, tokens, size_px, color=TEXT):
    for text, bold, code in tokens:
        run = para.add_run()
        run.text = text
        if code:
            set_run(run, size_px * 0.85, CODE_PINK, bold, mono=True)
        else:
            # 本文中の**のみアクセント色。見出しや白文字ではベース色を維持する
            run_color = PINK if bold and color is TEXT else color
            set_run(run, size_px, run_color, bold)


def set_gradient(shape):
    shape.fill.solid()
    spPr = shape._element.spPr
    old = spPr.find(qn("a:solidFill"))
    spPr.replace(old, parse_xml(GRAD_XML))
    shape.line.fill.background()
    shape.shadow.inherit = False


def set_bullet(para, ordered=False, color=PINK, indent_px=40):
    pPr = para._p.get_or_add_pPr()
    indent = round(indent_px * PX)
    pPr.set("marL", str(indent * (para.level + 1)))
    pPr.set("indent", str(-indent))
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    pPr.append(parse_xml(f'<a:buClr {ns}><a:srgbClr val="{color}"/></a:buClr>'))
    pPr.append(parse_xml(f'<a:buFont {ns} typeface="Arial"/>'))
    if ordered:
        pPr.append(parse_xml(f'<a:buAutoNum {ns} type="arabicPeriod"/>'))
    else:
        pPr.append(parse_xml(f'<a:buChar {ns} char="•"/>'))


def no_bullet(para):
    pPr = para._p.get_or_add_pPr()
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    pPr.append(parse_xml(f"<a:buNone {ns}/>"))


def add_textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        tcPr.append(parse_xml(
            f'<{tag} {ns} w="9525" cap="flat">'
            f'<a:solidFill><a:srgbClr val="DDDDE2"/></a:solidFill></{tag}>'
        ))


# ------------------------------------------------------------ slide build


def add_chrome(slide, header, page_no, dark=False):
    """全スライド共通のヘッダーとページ番号。"""
    if header:
        tf = add_textbox(slide, PAD_X, 20, CONTENT_W, 24)
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = header
        set_run(run, 15, RGBColor(0x73, 0x73, 0x7B) if dark else FAINT,
                bold=True, spacing_pt=2.5)
    tf = add_textbox(slide, SLIDE_W - 140, SLIDE_H - 46, 100, 28)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    run = tf.paragraphs[0].add_run()
    run.text = str(page_no)
    set_run(run, 18, MUTED)


def set_background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def lead_h1_height(lines):
    """leadのh1ブロックの高さ(px)。高さ見積もりと描画で共用する。

    2行以上あるとき、1行目はトピック番号(小さいラベル)なので44px行より低い。
    20px * line_spacing 1.3 + space_after 8pt = 26 + 10.67 ≒ 37px。
    """
    return (len(lines) - 1) * 62 + 37 if len(lines) > 1 else 62


def build_lead(slide, blocks):
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(SLIDE_W), px(10))
    set_gradient(strip)
    # コンテンツ量からおおまかに垂直センタリングする
    est = sum(lead_h1_height(b[1]) + 36 if b[0] == "h1" else 60 for b in blocks)
    y = max(120, (SLIDE_H - est) / 2 - 30)
    for block in blocks:
        if block[0] == "h1":
            lines = block[1]
            tf = add_textbox(slide, 90, y, SLIDE_W - 180, lead_h1_height(lines))
            for j, line in enumerate(lines):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.line_spacing = 1.3
                # 1行目はトピック番号。タイトルと同じ44px白では見分けがつかないので、
                # 小さいラベル(kicker)として扱う
                if j == 0 and len(lines) > 1:
                    p.space_after = Pt(8)
                    run = p.add_run()
                    run.text = plain_text(line)
                    set_run(run, 20, LEAD_SUB, spacing_pt=3.0)
                else:
                    fill_runs(p, [(plain_text(line), True, False)], 44, WHITE)
            y += lead_h1_height(lines) + 36
        elif block[0] == "p":
            tf = add_textbox(slide, 90, y, SLIDE_W - 180, 40)
            run = tf.paragraphs[0].add_run()
            run.text = plain_text(" ".join(block[1])).upper()
            set_run(run, 22, LEAD_SUB, bold=True, spacing_pt=3.0)
            y += 60


def table_geometry(rows, s):
    """(列数, 列幅リスト, 行高リスト) を返す。高さ見積もりと描画で共用する。"""
    n_cols = max(len(r) for r in rows)
    col_w = []
    for c in range(n_cols):
        longest = max(
            (text_units(plain_text(r[c])) for r in rows if c < len(r)), default=4)
        col_w.append(min(longest * 24 * s + 60 * s, 500 * s))
    fit = min(1.0, CONTENT_W / sum(col_w))
    col_w = [round(w * fit) for w in col_w]
    row_h = [
        round(max(est_lines(plain_text(c), 24 * s, col_w[j] - 40 * s)
                  for j, c in enumerate(r)) * 38 * s + 22 * s)
        for r in rows
    ]
    return n_cols, col_w, row_h


def block_height(block, lesson_dir, s=1.0):
    """コンテンツブロックの推定高さ(px)。倍率sは文字サイズに連動する。"""
    kind = block[0]
    if kind == "list":
        h = 0
        for level, _, text in block[1]:
            avail = CONTENT_W - 40 * s * (level + 1)
            h += est_lines(plain_text(text), 26 * s, avail) * 45 * s
        return h + 8 * s
    if kind == "p":
        return sum(est_lines(plain_text(l), 26 * s, CONTENT_W)
                   for l in block[1]) * 45 * s
    if kind == "code":
        return len(block[2]) * 33 * s + 46 * s
    if kind == "img":
        path = lesson_dir / block[1]
        png = path.with_suffix(".diagram.png") if path.suffix == ".svg" else path
        with Image.open(png) as im:
            iw, ih = im.size
        w = min(round(block[2] * s), CONTENT_W)
        return round(w * ih / iw) + 10
    if kind == "table":
        return sum(table_geometry(block[1], s)[2])
    return 40 * s


def render_block(slide, block, y, lesson_dir, s=1.0):
    kind = block[0]
    if kind == "list":
        items = block[1]
        h = block_height(block, lesson_dir, s)
        tf = add_textbox(slide, PAD_X, y, CONTENT_W, h)
        for j, (level, ordered, text) in enumerate(items):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.level = level
            p.line_spacing = 1.6
            p.space_after = Pt(4 * s)
            set_bullet(p, ordered, indent_px=40 * s)
            fill_runs(p, parse_inline(text), 26 * s)
        return y + h

    if kind == "p":
        h = block_height(block, lesson_dir, s)
        tf = add_textbox(slide, PAD_X, y, CONTENT_W, h)
        for j, line in enumerate(block[1]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.line_spacing = 1.6
            fill_runs(p, parse_inline(line), 26 * s)
        return y + h

    if kind == "code":
        h = block_height(block, lesson_dir, s)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, px(PAD_X), px(y), px(CONTENT_W), px(h))
        shape.adjustments[0] = min(0.5, 16 / h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = BG_SOFT
        shape.line.color.rgb = BORDER
        shape.line.width = Pt(0.75)
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = px(29 * s)
        tf.margin_top = tf.margin_bottom = px(20 * s)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for j, line_runs in enumerate(highlight_lines(block[2], block[1])):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.line_spacing = 1.5
            p.alignment = PP_ALIGN.LEFT  # オートシェイプの既定は中央寄せ
            for text, color in line_runs:
                run = p.add_run()
                run.text = text
                set_run(run, 22 * s, color, mono=True)
        return y + h

    if kind == "img":
        path = lesson_dir / block[1]
        png = path.with_suffix(".diagram.png") if path.suffix == ".svg" else path
        if not png.exists():
            raise FileNotFoundError(
                f"{png} がありません。先に python scripts/diagram_export.py を実行してください。")
        with Image.open(png) as im:
            iw, ih = im.size
        w = min(round(block[2] * s), CONTENT_W)
        h = round(w * ih / iw)
        slide.shapes.add_picture(str(png), px(PAD_X + (CONTENT_W - w) / 2), px(y),
                                 px(w), px(h))
        return y + h + 10

    if kind == "table":
        rows = block[1]
        n_cols, col_w, row_h = table_geometry(rows, s)
        total_w = sum(col_w)
        shape = slide.shapes.add_table(
            len(rows), n_cols,
            px(PAD_X + (CONTENT_W - total_w) / 2), px(y), px(total_w),
            px(sum(row_h)))
        table = shape.table
        table.first_row = False
        table.horz_banding = False
        for j, w in enumerate(col_w):
            table.columns[j].width = px(w)
        for ri, row in enumerate(rows):
            table.rows[ri].height = px(row_h[ri])
            for ci in range(n_cols):
                cell = table.cell(ri, ci)
                cell.margin_left = cell.margin_right = px(20 * s)
                cell.margin_top = cell.margin_bottom = px(10 * s)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                set_cell_borders(cell)
                cell.fill.solid()
                cell.fill.fore_color.rgb = INK if ri == 0 else WHITE
                text = row[ci] if ci < len(row) else ""
                p = cell.text_frame.paragraphs[0]
                if ri == 0:
                    p.alignment = PP_ALIGN.CENTER
                    fill_runs(p, [(plain_text(text), True, False)], 24 * s, WHITE)
                else:
                    fill_runs(p, parse_inline(text), 24 * s)
        return y + sum(row_h) + 10

    return y + 40 * s


def build_content(slide, blocks, lesson_dir, warn):
    # 内容量に応じて文字サイズの倍率を決め、スライド全体を使う。
    # 収まる範囲で最大1.45倍まで拡大し、あふれる場合は0.8倍まで縮小する
    avail = SLIDE_H - 60 - 45

    def total(s):
        h = 0.0
        for block in blocks:
            if block[0] in ("h1", "h2"):
                text = block[1] if block[0] == "h2" else " ".join(block[1])
                lines = est_lines(plain_text(text), 32.5 * s, CONTENT_W)
                h += lines * 44 * s + 10 * s + 4 + 28 * s
            else:
                h += block_height(block, lesson_dir, s) + 12 * s
        return h

    s = 1.45
    # コードは折り返さないので、最長行が枠内に収まる倍率を上限にする
    for block in blocks:
        if block[0] == "code":
            longest = max((mono_units(l) for l in block[2]), default=0)
            if longest:
                s = min(s, (CONTENT_W - 10) / (longest * 22 + 58))
    while s > 0.8 and total(s) > avail:
        s = round(s - 0.05, 2)

    y = 60
    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2"):
            text = block[1] if kind == "h2" else " ".join(block[1])
            lines = est_lines(plain_text(text), 32.5 * s, CONTENT_W)
            tf = add_textbox(slide, PAD_X, y, CONTENT_W, lines * 44 * s)
            p = tf.paragraphs[0]
            p.line_spacing = 1.3
            fill_runs(p, parse_inline(text), 32.5 * s, INK)
            y += lines * 44 * s + 10 * s
            bar = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, px(PAD_X), px(y), px(72), px(4))
            bar.adjustments[0] = 0.5
            set_gradient(bar)
            y += 4 + 28 * s
        else:
            y = render_block(slide, block, y + 6 * s, lesson_dir, s) + 6 * s
    if y > SLIDE_H - 30:
        warn(f"コンテンツがはみ出している可能性 (推定 {round(y)}px, 倍率 {s})")


def build_deck(md_path):
    md_path = Path(md_path)
    lesson_dir = md_path.parent
    header, body = parse_front_matter(md_path.read_text(encoding="utf-8"))
    slides_md = parse_slides(body)

    prs = Presentation()
    prs.slide_width = px(SLIDE_W)
    prs.slide_height = px(SLIDE_H)
    blank = prs.slide_layouts[6]
    warnings = []

    for idx, sl in enumerate(slides_md, start=1):
        slide = prs.slides.add_slide(blank)
        cls = sl["cls"]
        if cls == "lead":
            set_background(slide, LEAD_BG)
            add_chrome(slide, header, idx, dark=True)
            build_lead(slide, sl["blocks"])
        else:
            set_background(slide, BG_SOFT if cls == "summary" else WHITE)
            add_chrome(slide, header, idx)
            build_content(slide, sl["blocks"], lesson_dir,
                          lambda msg, i=idx: warnings.append(f"  slide {i}: {msg}"))
        if sl["note"]:
            slide.notes_slide.notes_text_frame.text = sl["note"]

    out = lesson_dir / "slides.pptx"
    prs.save(out)
    return out, warnings


def main():
    paths = sys.argv[1:]
    if not paths:
        print("usage: python scripts/build_pptx.py <slides.md> ...", file=sys.stderr)
        sys.exit(1)
    failed = False
    for i, p in enumerate(paths, start=1):
        try:
            out, warnings = build_deck(p)
            print(f"[{i}/{len(paths)}] {out} OK")
            for w in warnings:
                print(w)
        except Exception as e:  # noqa: BLE001 - 1件の失敗で全体を止めない
            failed = True
            print(f"[{i}/{len(paths)}] {p} 失敗: {e}", file=sys.stderr)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
